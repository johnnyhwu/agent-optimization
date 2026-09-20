"""Design system for the Skill Studio deck.

Values come from github.com/Gabberflast/academic-pptx-skill (slide_patterns.md):
16:9 at 10 x 5.625in, Arial throughout, white content slides, navy titles,
action title 26pt, body >= 20pt, inline labels 16pt, citations 13pt, 0.5in margins.

Everything in build_pptx.py composes slides out of these helpers, so the rules
live in one place instead of being retyped per slide.
"""
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.oxml import parse_xml
from PIL import Image

# ---------------------------------------------------------------- palette --
NAVY    = RGBColor(0x1F, 0x4E, 0x79)   # titles, dark slide backgrounds
ACCENT  = RGBColor(0x2E, 0x75, 0xB6)   # section headers, highlights
BODY    = RGBColor(0x2D, 0x2D, 0x2D)
MUTED   = RGBColor(0x77, 0x77, 0x77)
RULE    = RGBColor(0xCC, 0xCC, 0xCC)
WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
CALLOUT = RGBColor(0xEB, 0xF3, 0xFA)   # light blue callout fill
WARM    = RGBColor(0xFF, 0xF2, 0xCC)   # callout, warm
PALE    = RGBColor(0xF5, 0xF7, 0xFA)   # panel fill
LIGHT   = RGBColor(0xBD, 0xD7, 0xEE)   # text on navy
DIM     = RGBColor(0xB8, 0xBC, 0xC2)   # unlit pipeline stage

FONT = "Arial"
MONO = "Consolas"

# ------------------------------------------------------------------ sizes --
T_TITLE, T_SECTION, T_BODY, T_LABEL, T_CITE = 26, 22, 20, 16, 13

# geometry (inches)
SW, SH   = 10.0, 5.625
MARGIN   = 0.5
TITLE_Y  = 0.24
TITLE_H  = 0.98
RULE_Y   = 1.28
BODY_Y   = 1.48
BODY_H   = 3.55
CITE_Y   = 5.00
CW       = SW - 2 * MARGIN          # 9.0in content width


def _srgb(color):
    return '<a:solidFill><a:srgbClr val="%s"/></a:solidFill>' % f"{color}"


def blank(prs, bg=None):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    if bg is not None:
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = bg
    return s


def textbox(slide, x, y, w, h, *, anchor=MSO_ANCHOR.TOP, wrap=True):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return box


def para(tf, text, *, size=T_BODY, color=BODY, bold=False, italic=False,
         align=PP_ALIGN.LEFT, space_after=6, space_before=0, line=1.15,
         font=FONT, first=False):
    """Add a paragraph. `text` may be a list of (text, {overrides}) run tuples."""
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_after = Pt(space_after)
    p.space_before = Pt(space_before)
    p.line_spacing = line
    runs = text if isinstance(text, list) else [(text, {})]
    for chunk, over in runs:
        r = p.add_run()
        r.text = chunk
        f = r.font
        f.name = over.get("font", font)
        f.size = Pt(over.get("size", size))
        f.bold = over.get("bold", bold)
        f.italic = over.get("italic", italic)
        f.color.rgb = over.get("color", color)
    return p


def bullet(p, char="•", marL=0.26, hang=0.26):
    """Real hanging-indent bullet: wrapped lines align under the text, not the glyph."""
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(Emu(Inches(marL))))
    pPr.set("indent", str(-Emu(Inches(hang))))
    pPr.append(parse_xml('<a:buFont xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" typeface="Arial"/>'))
    pPr.append(parse_xml('<a:buChar xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" char="%s"/>' % char))
    return p


def action_title(slide, text, *, size=T_TITLE, color=NAVY, y=TITLE_Y, h=TITLE_H):
    """A complete sentence stating the takeaway — never a topic label."""
    box = textbox(slide, MARGIN, y, CW, h, anchor=MSO_ANCHOR.BOTTOM)
    box.name = "action-title"          # excluded from the body word count
    para(box.text_frame, text, size=size, color=color, bold=True,
         line=1.08, space_after=0, first=True)
    return box


def rule(slide, y=RULE_Y, x=MARGIN, w=CW, h=0.025, color=RULE):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    bar.fill.solid(); bar.fill.fore_color.rgb = color
    bar.line.fill.background(); bar.shadow.inherit = False
    return bar


def cite(slide, text, *, y=CITE_Y, x=MARGIN, w=CW, align=PP_ALIGN.LEFT, h=0.54):
    box = textbox(slide, x, y, w, h)
    para(box.text_frame, text, size=T_CITE, color=MUTED, align=align,
         space_after=0, first=True)
    return box


def header(slide, x, y, w, text, *, color=ACCENT, size=T_SECTION):
    box = textbox(slide, x, y, w, 0.34)
    para(box.text_frame, text, size=size, color=color, bold=True,
         space_after=0, first=True)
    return box


def bullets(slide, x, y, w, h, items, *, size=T_BODY, color=BODY, space_after=9, line=1.15):
    """items: str, or (text, {overrides}) run-list, one entry per bullet."""
    box = textbox(slide, x, y, w, h)
    tf = box.text_frame
    for i, item in enumerate(items):
        p = para(tf, item, size=size, color=color, space_after=space_after,
                 line=line, first=(i == 0))
        bullet(p)
    return box


def panel(slide, x, y, w, h, *, fill=PALE, line_color=None, radius=True):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h))
    if radius:
        shape.adjustments[0] = 0.06
    shape.fill.solid(); shape.fill.fore_color.rgb = fill
    if line_color is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line_color
        shape.line.width = Pt(1)
    shape.shadow.inherit = False
    shape.text_frame.word_wrap = True
    return shape


def arrow(slide, x1, y1, x2, y2, *, color=MUTED, width=1.25, dash=False, head=True):
    """A straight connector. DrawingML wants prstDash before tailEnd, so order matters."""
    from pptx.enum.shapes import MSO_CONNECTOR
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                   Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    if dash:
        ln.append(parse_xml('<a:prstDash xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" val="dash"/>'))
    if head:
        ln.append(parse_xml('<a:tailEnd xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" type="triangle" w="sm" len="sm"/>'))
    c.shadow.inherit = False
    return c


def picture(slide, path, x, y, max_w, max_h, *, center_x=True):
    """Place an image inside a box, preserving aspect ratio."""
    iw, ih = Image.open(path).size
    scale = min(max_w / iw, max_h / ih)
    w, h = iw * scale, ih * scale
    px = x + (max_w - w) / 2 if center_x else x
    py = y + (max_h - h) / 2
    return slide.shapes.add_picture(path, Inches(px), Inches(py), Inches(w), Inches(h))


def notes(slide, text):
    """Presenter crib sheet — Traditional Chinese, never projected."""
    slide.notes_slide.notes_text_frame.text = text.strip()
