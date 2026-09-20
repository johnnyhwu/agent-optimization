"""The three diagrams, drawn as native PowerPoint shapes.

Native shapes rather than images: the whole reason this deck ships as .pptx is
that the boxes can be dragged and the words retyped without regenerating anything.
"""
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from deck_theme import (NAVY, ACCENT, BODY, MUTED, RULE, WHITE, CALLOUT, PALE, DIM,
                        MARGIN, CW, textbox, para, panel, arrow)


def _badge(slide, cx, cy, n, *, color=ACCENT, d=0.24):
    o = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - d / 2), Inches(cy - d / 2),
                               Inches(d), Inches(d))
    o.fill.solid(); o.fill.fore_color.rgb = color
    o.line.fill.background(); o.shadow.inherit = False
    tf = o.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    para(tf, str(n), size=11, color=WHITE, bold=True, align=PP_ALIGN.CENTER,
         space_after=0, first=True)
    return o


# ------------------------------------------------------------------- loop --
LOOP_STEPS = [
    (0.55, 1.42, "1", "Configure it, run an eval", "the config is gone when the run ends"),
    (6.85, 1.42, "2", "Read the trace, find the step", "a dozen spans, opened by hand"),
    (6.85, 3.26, "3", "Edit the skill, test the hunch", "only one LLM call is testable"),
    (0.55, 3.26, "4", "A new question occurs to you", "keeping it means copying it by hand"),
]


def draw_loop(slide):
    for x, y, n, title, breaks in LOOP_STEPS:
        box = panel(slide, x, y, 2.60, 1.12, fill=WHITE, line_color=RULE)
        _badge(slide, x + 0.22, y + 0.25, n)
        t = textbox(slide, x + 0.42, y + 0.12, 2.04, 0.54)
        para(t.text_frame, title, size=13.5, color=BODY, bold=True, line=1.10,
             space_after=0, first=True)
        d = textbox(slide, x + 0.14, y + 0.64, 2.34, 0.44)
        para(d.text_frame, [("Breaks: ", {"color": ACCENT, "bold": True}), (breaks, {})],
             size=11.5, color=MUTED, line=1.12, space_after=0, first=True)

    arrow(slide, 3.25, 1.98, 6.75, 1.98)          # 1 -> 2
    arrow(slide, 8.15, 2.62, 8.15, 3.16)          # 2 -> 3
    arrow(slide, 6.75, 3.82, 3.25, 3.82)          # 3 -> 4
    arrow(slide, 1.85, 3.16, 1.85, 2.62)          # 4 -> 1

    c = textbox(slide, 3.45, 2.30, 3.10, 1.24, anchor=MSO_ANCHOR.MIDDLE)
    tf = c.text_frame
    para(tf, "and the loop itself takes", size=12, color=MUTED,
         align=PP_ALIGN.CENTER, space_after=2, first=True)
    para(tf, "3 days – 2 weeks", size=22, color=NAVY, bold=True,
         align=PP_ALIGN.CENTER, space_after=2)
    para(tf, "to take a new topic from 0% to above 90%", size=12, color=MUTED,
         align=PP_ALIGN.CENTER, space_after=0)


# ----------------------------------------------------------- architecture --
def draw_architecture(slide):
    panel(slide, 0.50, 2.36, 0.95, 0.60, fill=WHITE, line_color=RULE)
    b = textbox(slide, 0.50, 2.48, 0.95, 0.40)
    para(b.text_frame, "Browser", size=11.5, color=BODY, bold=True,
         align=PP_ALIGN.CENTER, space_after=0, first=True)
    arrow(slide, 1.50, 2.66, 1.62, 2.66)

    panel(slide, 1.66, 1.44, 4.58, 3.14, fill=WHITE, line_color=NAVY, radius=False)
    lab = textbox(slide, 1.80, 1.52, 2.0, 0.26)
    para(lab.text_frame, "Skill Studio", size=12, color=NAVY, bold=True, space_after=0, first=True)
    stack = textbox(slide, 4.10, 1.53, 2.0, 0.24)
    para(stack.text_frame, "FastAPI · Postgres", size=10, color=MUTED,
         align=PP_ALIGN.RIGHT, space_after=0, first=True)

    for i, name in enumerate(("Evaluation", "Playground", "Optimize")):
        x = 1.80 + i * 1.48
        p = panel(slide, x, 1.82, 1.40, 0.34, fill=CALLOUT, line_color=ACCENT)
        t = textbox(slide, x, 1.88, 1.40, 0.24)
        para(t.text_frame, name, size=11.5, color=ACCENT, bold=True,
             align=PP_ALIGN.CENTER, space_after=0, first=True)

    panel(slide, 1.80, 2.24, 4.30, 0.30, fill=PALE, line_color=RULE)
    o = textbox(slide, 1.80, 2.30, 4.30, 0.24)
    para(o.text_frame, "Orchestrator — background task, live per-question updates",
         size=10.5, color=BODY, align=PP_ALIGN.CENTER, space_after=0, first=True)

    s = textbox(slide, 1.80, 2.60, 4.30, 0.20)
    para(s.text_frame, "Seven swappable seams — fake and real, all fake by default",
         size=9.5, color=MUTED, space_after=0, first=True)
    seams = ["Agent", "Judge", "Trace", "Diagnosis", "Synthesis", "Workspace", "Optimizer"]
    for i, name in enumerate(seams):
        col, row = i % 4, i // 4
        x = 1.80 + col * 1.09
        y = 2.80 + row * 0.33
        panel(slide, x, y, 1.03, 0.28, fill=WHITE, line_color=RULE)
        t = textbox(slide, x, 2.845 + row * 0.33, 1.03, 0.22)
        para(t.text_frame, name, size=10, color=BODY, align=PP_ALIGN.CENTER,
             space_after=0, first=True)

    panel(slide, 1.80, 3.46, 4.30, 1.02, fill=PALE, line_color=RULE)
    d = textbox(slide, 1.92, 3.54, 4.06, 0.88)
    tf = d.text_frame
    para(tf, "Postgres — what Langfuse cannot express", size=11, color=NAVY,
         bold=True, space_after=3, first=True)
    para(tf, "eval sets · stable question ids · runs and their full config",
         size=10, color=BODY, space_after=2)
    para(tf, "verdicts · diagnoses · every step of an optimization run",
         size=10, color=BODY, space_after=2)
    para(tf, "span input / output / tokens are never copied in here", size=10, color=MUTED,
         space_after=0)

    ext = [(1.44, 0.58, "Agent Server", "OpenAI chat completions + two fields"),
           (2.54, 0.78, "Langfuse", "the only home of traces and spans;\nread live, never copied into our DB"),
           (3.42, 0.58, "LLM endpoints", "judge · diagnosis · synthesis · optimizer"),
           (4.04, 0.54, "Keycloak", "OIDC, optional")]
    for y, h, name, sub in ext:
        panel(slide, 6.52, y, 2.98, h, fill=WHITE, line_color=RULE)
        t = textbox(slide, 6.66, y + 0.07, 2.72, h - 0.10)
        tf = t.text_frame
        para(tf, name, size=11.5, color=BODY, bold=True, space_after=2, first=True)
        for line in sub.split("\n"):
            para(tf, line, size=9.5, color=MUTED, space_after=1, line=1.1)

    arrow(slide, 6.26, 2.90, 6.48, 1.73)
    arrow(slide, 6.26, 2.94, 6.48, 2.93)
    arrow(slide, 6.26, 3.00, 6.48, 3.71)
    arrow(slide, 6.26, 3.10, 6.48, 4.31)

    arrow(slide, 6.70, 2.06, 6.70, 2.50, color=ACCENT, dash=True)
    a = textbox(slide, 6.88, 2.06, 2.60, 0.46)
    tf = a.text_frame
    para(tf, "correlation id", size=11, color=ACCENT, bold=True, space_after=1, first=True)
    para(tf, "the agent reuses it as its trace id",
         size=9.5, color=ACCENT, space_after=0, line=1.1)


# --------------------------------------------------------------- pipeline --
STAGES = [("Rollout · split", 1), ("Reflection", 2), ("Hierarchical merge", 3),
          ("Rank · edit budget", 4), ("Validation gate", 5)]
_BW, _GAP, _Y, _BH = 1.704, 0.12, 1.44, 0.88


def stage_x(i):
    return MARGIN + i * (_BW + _GAP)


def draw_pipeline(slide, lit=None):
    """lit=None lights every stage; lit=n dims all but stage n."""
    for i, (name, n) in enumerate(STAGES):
        on = lit is None or lit == n
        x = stage_x(i)
        edge = ACCENT if on else DIM
        panel(slide, x, _Y, _BW, _BH, fill=CALLOUT if on else WHITE, line_color=edge)
        _badge(slide, x + 0.22, _Y + 0.22, n, color=ACCENT if on else DIM)
        t = textbox(slide, x + 0.06, _Y + 0.36, _BW - 0.12, 0.48)
        para(t.text_frame, name, size=12, color=BODY if on else DIM, bold=True,
             align=PP_ALIGN.CENTER, line=1.12, space_after=0, first=True)
        if i < 4:
            arrow(slide, x + _BW + 0.01, _Y + _BH / 2, x + _BW + _GAP - 0.01, _Y + _BH / 2,
                  color=MUTED if lit is None else DIM)

    out = textbox(slide, 5.60, _Y + _BH + 0.04, 3.90, 0.22)
    para(out.text_frame, [("pass → best_skill.md", {"color": ACCENT, "bold": True}),
                          ("    fail → step discarded", {"color": MUTED})],
         size=11, align=PP_ALIGN.RIGHT, space_after=0, first=True)

    fb_y = _Y + _BH + 0.56
    x5, x2 = stage_x(4) + _BW / 2, stage_x(1) + _BW / 2
    arrow(slide, x5, _Y + _BH + 0.30, x5, fb_y, color=DIM, width=1, head=False)
    arrow(slide, x5, fb_y, x2, fb_y, color=DIM, width=1, dash=True, head=False)
    arrow(slide, x2, fb_y, x2, _Y + _BH + 0.02, color=DIM, width=1)
    lab = textbox(slide, x2 + 0.20, fb_y - 0.24, x5 - x2 - 0.40, 0.22)
    para(lab.text_frame, "rejected edits feed the next step's analysis", size=10,
         color=MUTED, align=PP_ALIGN.CENTER, space_after=0, first=True)
